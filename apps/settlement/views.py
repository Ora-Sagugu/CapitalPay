"""清结算 — API Views (运营管理端)。"""
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP

from django.db.models import Sum, Q
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework.filters import SearchFilter, OrderingFilter

from apps.rbac.permissions import RequiresFeature
from apps.rbac.authentication import JWTAuthentication
from apps.core.excel_export import workbook_response
from .models import SettlementBatch, SettlementDetail, FeeShare, DifferenceWriteOff
from .serializers import (
    SettlementBatchSerializer, SettlementDetailSerializer,
    FeeShareSerializer, DifferenceWriteOffSerializer,
)

_MONEY = Decimal("0.01")


def _money_str(value) -> str:
    n = Decimal(str(value or 0))
    return format(n.quantize(_MONEY, rounding=ROUND_HALF_UP), "f")


class SettlementBatchViewSet(RequiresFeature, viewsets.ReadOnlyModelViewSet):
    """清算批次查询 — 运营管理端。"""
    authentication_classes = [JWTAuthentication]
    feature_code = "feature:reports"
    queryset = SettlementBatch.objects.filter(
        is_deleted=False
    ).select_related("merchant")
    serializer_class = SettlementBatchSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["merchant__merchant_no", "status", "settle_date"]
    search_fields = ["batch_no", "merchant__merchant_name"]
    ordering_fields = ["settle_date", "settle_net_amount", "total_amount"]

    @action(detail=True, methods=["post"], url_path="approve")
    def approve(self, request, pk=None):
        """审批清算批次 — PENDING → 触发资金划拨。"""
        from .engine.distributor import FundsDistributor
        batch = self.get_object()
        if batch.status != SettlementBatch.SettleStatus.PENDING:
            return Response({"detail": "Only a pending settlement batch may be approved"}, status=status.HTTP_400_BAD_REQUEST)
        distributor = FundsDistributor()
        batch = distributor.distribute(batch)
        return Response(SettlementBatchSerializer(batch).data)

    @action(detail=True, methods=["post", "get"], url_path="export")
    def export_batch(self, request, pk=None):
        """导出单个清算批次明细为 Excel。"""
        batch = self.get_object()
        details = list(
            SettlementDetail.objects.filter(batch=batch, is_deleted=False).order_by("created_at")
        )
        wb = Workbook()
        ws = wb.active
        ws.title = "清算明细"
        headers = ["序号", "订单号", "订单金额", "手续费", "结算金额"]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        for i, d in enumerate(details, 1):
            ws.cell(row=i + 1, column=1, value=i)
            ws.cell(row=i + 1, column=2, value=d.order_no)
            ws.cell(row=i + 1, column=3, value=float(d.amount))
            ws.cell(row=i + 1, column=4, value=float(d.fee))
            ws.cell(row=i + 1, column=5, value=float(d.settle_amount))

        return workbook_response(wb, f"settle_batch_{batch.batch_no}.xlsx")


class SettlementDetailViewSet(RequiresFeature, viewsets.ReadOnlyModelViewSet):
    feature_code = "feature:reports"
    """清算明细查询 — 运营管理端。"""
    queryset = SettlementDetail.objects.filter(
        is_deleted=False
    ).select_related("batch__merchant")
    serializer_class = SettlementDetailSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["batch__batch_no", "batch__merchant__merchant_no", "batch__settle_date"]
    search_fields = ["order_no"]
    ordering_fields = ["amount", "fee", "settle_amount", "created_at"]


class FeeShareViewSet(RequiresFeature, viewsets.ReadOnlyModelViewSet):
    feature_code = "feature:agent_fees"
    """手续费分润查询 — 分润结算管理核心 API。"""
    queryset = FeeShare.objects.filter(
        is_deleted=False
    ).select_related("payment_order__merchant__agent", "agent")
    serializer_class = FeeShareSerializer
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["agent__agent_no", "agent_name", "merchant_name", "bank_channel_name"]
    search_fields = ["order_no", "agent_name", "merchant_name", "bank_channel_name",
                     "payment_order__beneficiary_name", "payment_order__beneficiary_bank"]
    ordering_fields = [
        "amount", "total_fee", "channel_fee", "platform_fee", "agent_fee", "created_at",
    ]
    ordering = ["-created_at"]

    def get_queryset(self):
        qs = super().get_queryset()
        # 日期范围筛选（基于汇款订单创建时间）
        start_date = self.request.query_params.get("start_date")
        end_date = self.request.query_params.get("end_date")
        if start_date:
            qs = qs.filter(payment_order__created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(payment_order__created_at__date__lte=end_date)
        # 关键词搜索（支持多字段模糊匹配）
        keyword = self.request.query_params.get("keyword")
        if keyword:
            qs = qs.filter(
                Q(order_no__icontains=keyword) |
                Q(agent_name__icontains=keyword) |
                Q(merchant_name__icontains=keyword) |
                Q(bank_channel_name__icontains=keyword) |
                Q(payment_order__beneficiary_name__icontains=keyword) |
                Q(payment_order__beneficiary_bank__icontains=keyword)
            )
        return qs

    @action(detail=False, methods=["get", "post"], url_path="stats")
    def stats(self, request):
        """分润汇总统计。

        请求体:
            { start_date, end_date, agent_name, merchant_name, bank_channel_name }

        返回:
            { total_orders, total_amount, total_fee, total_channel_fee,
              total_platform_fee, total_agent_fee }
        """
        qs = self.get_queryset()
        start_date = request.data.get("start_date") or request.query_params.get("start_date")
        end_date = request.data.get("end_date") or request.query_params.get("end_date")
        agent_name = request.data.get("agent_name")
        merchant_name = request.data.get("merchant_name")
        bank_channel_name = request.data.get("bank_channel_name")

        if start_date:
            qs = qs.filter(payment_order__created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(payment_order__created_at__date__lte=end_date)
        if agent_name:
            qs = qs.filter(agent_name__icontains=agent_name)
        if merchant_name:
            qs = qs.filter(merchant_name__icontains=merchant_name)
        if bank_channel_name:
            qs = qs.filter(bank_channel_name__icontains=bank_channel_name)

        aggregates = qs.aggregate(
            total_amount=Sum("amount"),
            total_fee=Sum("total_fee"),
            total_channel_fee=Sum("channel_fee"),
            total_platform_fee=Sum("platform_fee"),
            total_agent_fee=Sum("agent_fee"),
        )
        return Response({
            "total_orders": qs.count(),
            "total_amount": _money_str(aggregates["total_amount"]),
            "total_fee": _money_str(aggregates["total_fee"]),
            "total_channel_fee": _money_str(aggregates["total_channel_fee"]),
            "total_platform_fee": _money_str(aggregates["total_platform_fee"]),
            "total_agent_fee": _money_str(aggregates["total_agent_fee"]),
        })

    @action(detail=False, methods=["post"], url_path="export")
    def export(self, request):
        """导出分润明细为 Excel。

        请求体:
            { start_date, end_date, agent_name, merchant_name, bank_channel_name }
        """
        qs = self.get_queryset()
        start_date = request.data.get("start_date") or request.query_params.get("start_date")
        end_date = request.data.get("end_date") or request.query_params.get("end_date")
        agent_name = request.data.get("agent_name")
        merchant_name = request.data.get("merchant_name")
        bank_channel_name = request.data.get("bank_channel_name")

        if start_date:
            qs = qs.filter(payment_order__created_at__date__gte=start_date)
        if end_date:
            qs = qs.filter(payment_order__created_at__date__lte=end_date)
        if agent_name:
            qs = qs.filter(agent_name__icontains=agent_name)
        if merchant_name:
            qs = qs.filter(merchant_name__icontains=merchant_name)
        if bank_channel_name:
            qs = qs.filter(bank_channel_name__icontains=bank_channel_name)

        records = list(qs)

        wb = Workbook()
        ws = wb.active
        ws.title = "分润结算明细"

        # ── 样式 ──
        header_font = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="E17055", end_color="E17055", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell_align = Alignment(horizontal="center", vertical="center")
        money_align = Alignment(horizontal="right", vertical="center")
        thin_border = Border(
            left=Side(style="thin", color="D0D0D0"),
            right=Side(style="thin", color="D0D0D0"),
            top=Side(style="thin", color="D0D0D0"),
            bottom=Side(style="thin", color="D0D0D0"),
        )
        stripe_fill = PatternFill(start_color="FFF5EE", end_color="FFF5EE", fill_type="solid")

        # ── 标题行 ──
        ws.merge_cells("A1:K1")
        title_cell = ws["A1"]
        title_cell.value = "手续费分润结算明细表"
        title_cell.font = Font(name="微软雅黑", size=14, bold=True, color="E17055")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 36

        # ── 汇总统计 ──
        total_amount = sum(r.amount for r in records) if records else 0
        total_fee = sum(r.total_fee for r in records) if records else 0
        total_channel = sum(r.channel_fee for r in records) if records else 0
        total_platform = sum(r.platform_fee for r in records) if records else 0
        total_agent = sum(r.agent_fee for r in records) if records else 0

        summary_headers = [
            "导出时间", f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "总订单数", str(len(records)),
            "交易总额", f"¥{total_amount:,.2f}",
            "手续费总额", f"¥{total_fee:,.2f}",
            "渠道手续费合计", f"¥{total_channel:,.2f}",
            "平台净收益合计", f"¥{total_platform:,.2f}",
            "代理商佣金合计", f"¥{total_agent:,.2f}",
        ]
        for col, val in enumerate(summary_headers, 1):
            c = ws.cell(row=3, column=col, value=val)
            c.font = Font(name="微软雅黑", size=10, bold=True if col % 2 == 1 else False,
                          color="333333")
            c.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[3].height = 28

        # ── 表头 ──
        headers = [
            "序号", "订单号", "清算日期", "商户名称", "代理商",
            "银行渠道", "订单金额", "手续费总额",
            "渠道手续费", "平台净收益", "代理商佣金",
        ]
        header_row = 5
        for col, h in enumerate(headers, 1):
            c = ws.cell(row=header_row, column=col, value=h)
            c.font = header_font
            c.fill = header_fill
            c.alignment = header_align
            c.border = thin_border
        ws.row_dimensions[header_row].height = 32

        # ── 数据行 ──
        col_widths = [6, 20, 14, 22, 18, 22, 16, 16, 16, 16, 16]
        for i, r in enumerate(records):
            row = header_row + 1 + i
            settle_date = ""
            try:
                po = r.payment_order
                if po:
                    dt = po.settled_at or po.created_at
                    settle_date = dt.strftime("%Y-%m-%d") if dt else ""
            except Exception:
                pass

            row_data = [
                i + 1,
                r.order_no,
                settle_date,
                r.merchant_name,
                r.agent_name or "-",
                r.bank_channel_name or "-",
                float(r.amount),
                float(r.total_fee),
                float(r.channel_fee),
                float(r.platform_fee),
                float(r.agent_fee),
            ]
            for col, val in enumerate(row_data, 1):
                c = ws.cell(row=row, column=col, value=val)
                c.font = Font(name="微软雅黑", size=10, color="333333")
                c.alignment = money_align if col >= 7 else cell_align
                c.border = thin_border
                if col >= 7 and isinstance(val, (int, float)):
                    c.number_format = '#,##0.00'
                # 条纹
                if i % 2 == 1:
                    c.fill = stripe_fill
            ws.row_dimensions[row].height = 24

        # ── 列宽 ──
        for col, w in enumerate(col_widths, 1):
            ws.column_dimensions[ws.cell(row=header_row, column=col).column_letter].width = w

        # ── 冻结窗格 ──
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

        filename = f"profit_share_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return workbook_response(wb, filename)


class DifferenceWriteOffViewSet(RequiresFeature, viewsets.ModelViewSet):
    feature_code = "feature:adjustments"
    ACTION_FEATURES = {
        "approve": "feature:adjustments.approve",
        "write_off": "feature:adjustments.approve",
    }
    """差异代销账管理 — 运营管理端。"""
    queryset = DifferenceWriteOff.objects.filter(is_deleted=False).select_related("merchant")
    serializer_class = DifferenceWriteOffSerializer
    lookup_field = "write_off_no"
    filter_backends = [DjangoFilterBackend, SearchFilter, OrderingFilter]
    filterset_fields = ["merchant__merchant_no", "status"]
    search_fields = ["write_off_no", "reason", "applied_by"]

    def perform_create(self, serializer):
        from apps.core.utils import generate_batch_no
        serializer.save(
            write_off_no=generate_batch_no("WOF"),
            applied_by=serializer.validated_data.get("applied_by") or "admin",
        )

    @action(detail=True, methods=["post"], url_path="approve")
    def approve(self, request, write_off_no=None):
        """审批差异代销账申请。"""
        from django.utils import timezone
        write_off = self.get_object()
        if write_off.status != DifferenceWriteOff.WriteOffStatus.PENDING:
            return Response({"message": "The current status does not permit approval"}, status=400)

        approved_by = request.data.get("approved_by", "system")
        write_off.status = DifferenceWriteOff.WriteOffStatus.APPROVED
        write_off.approved_by = approved_by
        write_off.approved_at = timezone.now()
        write_off.save()
        return Response({"message": "Authorisation recorded"})

    @action(detail=True, methods=["post"], url_path="write-off")
    def write_off(self, request, write_off_no=None):
        """核销入账。"""
        from decimal import Decimal
        from apps.account.models import NostroAccount
        from apps.account.services import AccountService

        write_off = self.get_object()
        if write_off.status != DifferenceWriteOff.WriteOffStatus.APPROVED:
            return Response({"message": "Only an approved record may be written off"}, status=400)
        account = NostroAccount.objects.filter(
            merchant=write_off.merchant, is_deleted=False
        ).first()
        if account:
            AccountService().update_balance(account, Decimal(str(write_off.amount)), is_credit=True)
        write_off.status = DifferenceWriteOff.WriteOffStatus.WRITTEN_OFF
        write_off.save(update_fields=["status", "updated_at"])
        return Response({"message": "The write-off has been posted"})
