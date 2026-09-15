"""Generate CapitalPay topology slides in the template visual style.

Creates docs/CapitalPay_Topology.pptx — three widescreen slides:
  1. Application Architecture (Network / Application Service / Basic Component)
  2. System Topology (clients, Nginx+Django, stores, external integrations)
  3. Infra Architecture Flow (four layers including Database)

Usage:
    python scripts/build_topology_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt


ROOT = Path(__file__).resolve().parent.parent
OUT_PATH = ROOT / "docs" / "CapitalPay_Topology.pptx"

# Widescreen 16:9 — same as Technical Solution Template.pptx
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Template palette
POWDER = RGBColor(0xB0, 0xE0, 0xE6)  # Network
SKY = RGBColor(0x87, 0xCE, 0xFA)  # Application Service
CORNFLOWER = RGBColor(0x64, 0x95, 0xED)  # Basic Component
NAVY = RGBColor(0x06, 0x48, 0x90)  # titles / connectors
INK = RGBColor(0x1D, 0x34, 0x44)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
SAND = RGBColor(0xEE, 0xC9, 0xAF)  # External
MINT = RGBColor(0xC5, 0xE1, 0xA5)  # Clients
LILAC = RGBColor(0xD1, 0xC4, 0xE9)  # Scheduled jobs
FOOTNOTE = RGBColor(0x5A, 0x6A, 0x72)
HEADER_BG = RGBColor(0x4C, 0xA8, 0xE8)


def _set_run(paragraph, text, size, bold, color, name="Calibri"):
    paragraph.text = text
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = name
    return paragraph.runs[0] if paragraph.runs else None


def _anchor_center(tf):
    tf.word_wrap = True
    tf.auto_size = None
    body_pr = tf._txBody.find(qn("a:bodyPr"))
    if body_pr is not None:
        body_pr.set("anchor", "ctr")
        body_pr.set("lIns", str(Emu(Inches(0.04))))
        body_pr.set("rIns", str(Emu(Inches(0.04))))
        body_pr.set("tIns", str(Emu(Inches(0.02))))
        body_pr.set("bIns", str(Emu(Inches(0.02))))


def _style_paragraph(paragraph, size, bold, color, name="Calibri"):
    paragraph.alignment = PP_ALIGN.CENTER
    for run in paragraph.runs:
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = name


def add_box(
    slide,
    left,
    top,
    width,
    height,
    text,
    fill,
    font_size=11,
    bold=True,
    font_color=INK,
    line=None,
    radius=0.18,
):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height
    )
    try:
        shape.adjustments[0] = radius
    except (IndexError, ValueError):
        pass
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)
    tf = shape.text_frame
    _anchor_center(tf)
    if not text:
        tf.text = ""
        return shape
    lines = text.split("\n")
    tf.text = lines[0]
    _style_paragraph(tf.paragraphs[0], font_size, bold, font_color)
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        _style_paragraph(p, font_size, bold, font_color)
    return shape


def add_plain_label(slide, left, top, width, height, text, size=13, bold=True, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    _set_run(p, text, size, bold, color)
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    return box


def add_left_label(slide, left, top, width, height, text, size=12):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    lines = text.split("\n")
    tf.text = lines[0]
    _style_paragraph(tf.paragraphs[0], size, True, NAVY)
    for line in lines[1:]:
        p = tf.add_paragraph()
        p.text = line
        _style_paragraph(p, size, True, NAVY)
    return box


def add_title(slide, text):
    add_plain_label(
        slide,
        Inches(0.3),
        Inches(0.08),
        Inches(12.7),
        Inches(0.36),
        text,
        size=18,
        bold=True,
        color=NAVY,
    )


def add_connector(slide, x1, y1, x2, y2, dashed=False, color=NAVY, width=1.5):
    cxn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2
    )
    cxn.line.color.rgb = color
    cxn.line.width = Pt(width)
    if dashed:
        from pptx.enum.dml import MSO_LINE_DASH_STYLE

        cxn.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    return cxn


def _mid_right(box):
    return box.left + box.width, box.top + box.height // 2


def _mid_left(box):
    return box.left, box.top + box.height // 2


def _mid_bottom(box):
    return box.left + box.width // 2, box.top + box.height


def _mid_top(box):
    return box.left + box.width // 2, box.top


def connect_lr(slide, a, b, dashed=False):
    x1, y1 = _mid_right(a)
    x2, y2 = _mid_left(b)
    add_connector(slide, x1, y1, x2, y2, dashed=dashed)


def connect_tb(slide, a, b, dashed=False):
    x1, y1 = _mid_bottom(a)
    x2, y2 = _mid_top(b)
    add_connector(slide, x1, y1, x2, y2, dashed=dashed)


# ── Slide 1: Application Architecture ────────────────────────────────


def build_slide1(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    add_title(slide, "CapitalPay — Application Architecture")

    label_x, label_w = Inches(0.18), Inches(1.72)
    content_x = Inches(1.98)
    content_right = Inches(13.12)
    content_w = content_right - content_x

    # Network
    net_y = Inches(0.52)
    add_left_label(slide, label_x, net_y, label_w, Inches(0.92), "Network")
    half = (content_w - Inches(0.1)) // 2
    add_box(slide, content_x, net_y, half, Inches(0.42), "Nginx", POWDER, 12)
    add_box(
        slide,
        content_x + half + Inches(0.1),
        net_y,
        half,
        Inches(0.42),
        "Gunicorn",
        POWDER,
        12,
    )
    add_box(
        slide,
        content_x,
        net_y + Inches(0.50),
        content_w,
        Inches(0.40),
        "Reverse Proxy  (Nginx → 127.0.0.1:1024, not a separate LB cluster)",
        POWDER,
        11,
        bold=False,
    )

    # Application Service — 5 columns
    app_top = Inches(1.58)
    app_bottom = Inches(5.78)
    add_left_label(
        slide,
        label_x,
        Inches(3.15),
        label_w,
        Inches(0.80),
        "Application\nService",
        size=12,
    )

    columns = [
        (
            "Payment / Remittance",
            [
                "Pre-order",
                "UIN matching",
                "Payment confirm",
                "Close order",
                "Refund",
                "Cashier",
                "Fund trace",
                "Disbursement",
            ],
        ),
        (
            "Merchant / Onboarding",
            [
                "Merchant lifecycle",
                "Ops KYC",
                "Customer onboarding",
                "Agent KYC",
                "Agent / PRN",
                "Virtual account",
                "Deposit",
            ],
        ),
        (
            "Routing / FX",
            [
                "Bank channel routing",
                "Bank gateway",
                "FX inquiry",
                "Partner banks",
            ],
        ),
        (
            "Compliance / Recon",
            [
                "OFAC / UN sanctions",
                "Sanction scan",
                "Daily recon",
                "Nostro check",
                "Adjustment",
            ],
        ),
        (
            "Settlement / Report",
            [
                "Daily settlement",
                "Fee share",
                "Analytics",
                "Daily reports",
                "Fund transfer",
            ],
        ),
    ]

    gap = Inches(0.10)
    n_cols = len(columns)
    col_w = int((content_w - gap * (n_cols - 1)) / n_cols)
    header_h = Inches(0.32)
    box_h = Inches(0.42)
    box_gap = Inches(0.06)

    for i, (header, items) in enumerate(columns):
        x = content_x + i * (col_w + gap)
        add_box(slide, x, app_top, col_w, header_h, header, HEADER_BG, 10, True, WHITE)
        y = app_top + header_h + Inches(0.08)
        for item in items:
            add_box(slide, x, y, col_w, box_h, item, SKY, 10)
            y += box_h + box_gap

    # Basic Component
    basic_y = Inches(6.00)
    add_left_label(
        slide, label_x, basic_y, label_w, Inches(1.10), "Basic\nComponent", size=12
    )
    basics = [
        "System Params",
        "RBAC / Role",
        "JWT + HMAC Auth",
        "Nostro A/C",
        "Audit Log",
        "Fernet Encryption",
        "Scheduled Jobs",
        "OpenAPI",
    ]
    b_gap = Inches(0.08)
    b_w = int((content_w - b_gap * 3) / 4)
    b_h = Inches(0.40)
    for i, name in enumerate(basics):
        row, col = divmod(i, 4)
        x = content_x + col * (b_w + b_gap)
        y = basic_y + row * (b_h + Inches(0.08))
        add_box(slide, x, y, b_w, b_h, name, CORNFLOWER, 10, True, WHITE)

    _ = app_bottom  # layout bound used as visual guide only


# ── Slide 2: System Topology ─────────────────────────────────────────


def build_slide2(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "CapitalPay — System Topology")
    add_plain_label(
        slide,
        Inches(0.3),
        Inches(0.42),
        Inches(12.7),
        Inches(0.28),
        "Django monolith behind Nginx  ·  modules are functions, not independently deployable services",
        size=11,
        bold=False,
        color=FOOTNOTE,
    )

    # Lane headers
    add_plain_label(slide, Inches(0.28), Inches(0.72), Inches(2.35), Inches(0.28), "Clients", 12)
    add_plain_label(slide, Inches(3.05), Inches(0.72), Inches(3.55), Inches(0.28), "Platform", 12)
    add_plain_label(slide, Inches(9.70), Inches(0.72), Inches(3.25), Inches(0.28), "External", 12)
    add_plain_label(slide, Inches(7.12), Inches(3.12), Inches(2.28), Inches(0.26), "Data", 12)

    clients = [
        "Merchant HMAC API",
        "Agent PRN API",
        "Ops Portal  Vue :1025",
        "Agent :1026 / Customer :1027",
        "Public Cashier",
    ]
    client_boxes = []
    cy = Inches(1.02)
    for title in clients:
        box = add_box(slide, Inches(0.28), cy, Inches(2.40), Inches(0.62), title, MINT, 11)
        client_boxes.append(box)
        cy += Inches(0.72)

    nginx = add_box(
        slide, Inches(3.15), Inches(2.40), Inches(1.70), Inches(0.85), "Nginx\n:80", POWDER, 13
    )
    django = add_box(
        slide,
        Inches(5.20),
        Inches(2.05),
        Inches(1.70),
        Inches(1.45),
        "Django\nGunicorn\n:1024",
        SKY,
        13,
    )
    cron = add_box(
        slide,
        Inches(5.20),
        Inches(3.68),
        Inches(1.70),
        Inches(0.62),
        "Cron\nscheduled jobs",
        LILAC,
        11,
    )

    data_panel = add_box(
        slide,
        Inches(7.12),
        Inches(3.42),
        Inches(2.28),
        Inches(1.55),
        "",
        RGBColor(0xE8, 0xF0, 0xFE),
        line=CORNFLOWER,
        radius=0.08,
    )
    pg = add_box(
        slide,
        Inches(7.25),
        Inches(3.50),
        Inches(2.02),
        Inches(0.42),
        "MySQL 8 (prod)",
        CORNFLOWER,
        11,
        True,
        WHITE,
    )
    media = add_box(
        slide,
        Inches(7.25),
        Inches(3.96),
        Inches(2.02),
        Inches(0.42),
        "Media / Static",
        CORNFLOWER,
        11,
        True,
        WHITE,
    )
    cache = add_box(
        slide,
        Inches(7.25),
        Inches(4.42),
        Inches(2.02),
        Inches(0.42),
        "LocMemCache nonce",
        CORNFLOWER,
        11,
        True,
        WHITE,
    )

    ext_panel = add_box(
        slide,
        Inches(9.68),
        Inches(1.02),
        Inches(3.38),
        Inches(3.90),
        "",
        RGBColor(0xFD, 0xF3, 0xE8),
        line=RGBColor(0xD4, 0xA5, 0x7A),
        radius=0.08,
    )
    externals = [
        "Bank Gateway  MOCK / REAL",
        "Recon fetch  MOCK / SFTP / API",
        "FX provider",
        "SMS",
        "OFAC SDN / UN lists",
        "Merchant notify_url",
    ]
    ext_boxes = []
    ey = Inches(1.14)
    for name in externals:
        box = add_box(slide, Inches(9.82), ey, Inches(3.10), Inches(0.50), name, SAND, 11)
        ext_boxes.append(box)
        ey += Inches(0.60)

    for box in client_boxes:
        connect_lr(slide, box, nginx)
    connect_lr(slide, nginx, django)
    connect_tb(slide, django, cron)
    connect_lr(slide, django, data_panel)
    connect_lr(slide, django, ext_panel, dashed=True)

    add_plain_label(
        slide,
        Inches(0.28),
        Inches(5.12),
        Inches(12.7),
        Inches(0.32),
        "Auth   ·   /api/v1/ HMAC-SHA256 (merchant + agent PRN)     "
        "/api/v1/admin/ JWT (ops)     /api/v1/user/ JWT (customer & agent portal)     "
        "/api/v1/cashier/ public order page",
        size=11,
        bold=False,
        color=INK,
    )
    add_plain_label(
        slide,
        Inches(0.28),
        Inches(5.44),
        Inches(12.7),
        Inches(0.32),
        "Jobs   ·   close_expired_orders  ·  retry_failed_notifications  ·  "
        "run_daily_reconciliation  ·  run_daily_settlement  ·  generate_daily_reports  ·  "
        "refresh_sanction_lists  ·  suspend_expired_licenses",
        size=10,
        bold=False,
        color=FOOTNOTE,
    )
    add_plain_label(
        slide,
        Inches(0.28),
        Inches(5.76),
        Inches(12.7),
        Inches(0.32),
        "Local MVP uses SQLite + LocMemCache (no Redis / Celery).  "
        "Bank, FX, SMS, and recon adapters default to MOCK until credentials are set.",
        size=10,
        bold=False,
        color=FOOTNOTE,
    )
    _ = (pg, media, cache, ext_boxes)


# ── Slide 3: Infra Architecture Flow ─────────────────────────────────


def build_slide3(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, "CapitalPay — Infra Architecture Flow")

    label_x, label_w = Inches(0.22), Inches(1.70)
    content_x = Inches(2.05)
    content_w = Inches(10.95)

    layers = [
        (
            "Network",
            POWDER,
            INK,
            Inches(0.95),
            [("Nginx", 2.2), ("Gunicorn", 2.2), ("Reverse Proxy", 2.6)],
            False,
        ),
        (
            "Application\nService",
            SKY,
            INK,
            Inches(2.05),
            [
                ("Payment /\nRemittance", 1.9),
                ("Merchant /\nOnboarding", 1.9),
                ("Routing / FX", 1.9),
                ("Compliance /\nRecon", 1.9),
                ("Settlement /\nReport", 1.9),
            ],
            False,
        ),
        (
            "Basic\nComponent",
            CORNFLOWER,
            WHITE,
            Inches(3.95),
            [
                ("RBAC / JWT", 1.30),
                ("HMAC Auth", 1.30),
                ("System Params", 1.45),
                ("Nostro A/C", 1.30),
                ("Audit Log", 1.20),
                ("Fernet", 1.10),
                ("Scheduled Jobs", 1.45),
                ("OpenAPI", 1.15),
            ],
            False,
        ),
        (
            "Database",
            NAVY,
            WHITE,
            Inches(4.90),
            [
                ("MySQL 8.0  (production)", 3.6),
                ("Media / Static files", 2.8),
            ],
            False,
        ),
    ]

    # Sub-items under Application Service
    app_details = [
        ["Pre-order", "Refund", "Cashier"],
        ["KYC", "Onboarding", "Agent / PRN"],
        ["Bank gateway", "FX", "Routing"],
        ["OFAC / UN", "Recon", "Adjustment"],
        ["Settlement", "Fee share", "Reports"],
    ]

    for label, color, font_c, y, items, _ in layers:
        add_left_label(slide, label_x, y, label_w, Inches(0.70), label, size=12)
        total_w = sum(Inches(w) for _, w in items) + Inches(0.10) * (len(items) - 1)
        # center the row in content area if shorter
        x = content_x
        extra = content_w - total_w
        if extra > 0 and label != "Application\nService":
            x = content_x + extra // 2
        h = Inches(0.58) if "Application" not in label else Inches(0.50)
        boxes = []
        for name, w_in in items:
            box = add_box(
                slide, x, y, Inches(w_in), h, name, color, 11, True, font_c
            )
            boxes.append(box)
            x += Inches(w_in) + Inches(0.10)

        if "Application" in label:
            dy = y + Inches(0.58)
            x = content_x
            for col, (header_w) in zip(app_details, items):
                w_in = header_w[1]
                sub_h = Inches(0.32)
                sy = dy
                for sub in col:
                    add_box(slide, x, sy, Inches(w_in), sub_h, sub, SKY, 9, False)
                    sy += sub_h + Inches(0.05)
                x += Inches(w_in) + Inches(0.10)

    add_plain_label(
        slide,
        Inches(0.30),
        Inches(5.75),
        Inches(12.7),
        Inches(0.70),
        "Local MVP  =  SQLite + LocMemCache   ·   no Redis, no Celery, no separate load-balancer.  "
        "Production  =  Nginx + Gunicorn (systemd) + MySQL 8.0.  "
        "Bank / FX / SMS / recon stay MOCK until credentials are configured.",
        size=11,
        bold=False,
        color=FOOTNOTE,
    )


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    build_slide1(prs)
    build_slide2(prs)
    build_slide3(prs)
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PATH))
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
