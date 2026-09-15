"""
Edit docs/CapitalPay Main Process.pptx in place only — never write a new PPT file.

Usage (from repo root):
  .venv\\Scripts\\python.exe scripts\\build_sow_ppt.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SHOTS = ROOT / "docs" / "sow_screenshots"
OUT = ROOT / "docs" / "CapitalPay Main Process.pptx"

# Widescreen 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

BG = RGBColor(0xFA, 0xFA, 0xF8)
TITLE = RGBColor(0x1F, 0x29, 0x37)
ACCENT = RGBColor(0xE8, 0x6A, 0x17)
MUTED = RGBColor(0x5B, 0x64, 0x72)
LINE = RGBColor(0xE5, 0xE7, 0xEB)


def set_slide_bg(slide, color: RGBColor = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_textbox(slide, left, top, width, height, text, *, size=18, bold=False, color=TITLE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align
    return box


def add_caption(slide, text: str) -> None:
    add_textbox(
        slide,
        Inches(0.6),
        Inches(6.85),
        Inches(12.1),
        Inches(0.5),
        text,
        size=13,
        color=MUTED,
        align=PP_ALIGN.LEFT,
    )


def add_title(slide, text: str) -> None:
    add_textbox(slide, Inches(0.6), Inches(0.28), Inches(12.1), Inches(0.5), text, size=26, bold=True, color=TITLE)
    # accent underline
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.6), Inches(0.78), Inches(1.2), Inches(0.05))
    shape.fill.solid()
    shape.fill.fore_color.rgb = ACCENT
    shape.line.fill.background()


def fit_image(slide, image_path: Path, *, top=Inches(1.0), max_w=Inches(12.1), max_h=Inches(5.6)) -> None:
    """Place screenshot centered under title, preserving aspect ratio."""
    from PIL import Image

    left = Inches(0.6)
    with Image.open(image_path) as im:
        iw, ih = im.size
    aspect = iw / ih
    box_aspect = max_w / max_h
    if aspect >= box_aspect:
        width = max_w
        height = Emu(int(max_w / aspect))
    else:
        height = max_h
        width = Emu(int(max_h * aspect))
    left = Inches(0.6) + (max_w - width) / 2
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def content_slide(prs, title: str, image_name: str | None, caption: str) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide)
    add_title(slide, title)
    if image_name:
        path = SHOTS / image_name
        if path.exists():
            fit_image(slide, path)
        else:
            add_textbox(slide, Inches(0.6), Inches(2.5), Inches(12), Inches(1), f"[Missing screenshot: {image_name}]", size=16, color=MUTED)
    add_caption(slide, caption)


def cover_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    # accent bar
    bar = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), SLIDE_H)
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()

    add_textbox(slide, Inches(1.0), Inches(2.2), Inches(11), Inches(0.6), "CapitalPay", size=40, bold=True, color=ACCENT)
    add_textbox(slide, Inches(1.0), Inches(3.0), Inches(11), Inches(0.7), "SOW — Main Process Overview", size=32, bold=True, color=TITLE)
    add_textbox(
        slide,
        Inches(1.0),
        Inches(3.9),
        Inches(11),
        Inches(0.8),
        "Registration & login · Fee configuration · Remittance · PRN · Confirm payment · Smart payout routing · Status tracking · Account ledger",
        size=16,
        color=MUTED,
    )
    add_textbox(slide, Inches(1.0), Inches(6.6), Inches(11), Inches(0.4), "Demo portals: Ops :1025 · Agent :1026 · Customer :1027", size=12, color=MUTED)


def process_map_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "End-to-End Main Process")

    steps = [
        ("1. Register\n& Login", "Agent / Customer"),
        ("2. Fee\nConfig", "Remittance &\nAgent Fee"),
        ("3. Remittance\nApply", "Quote &\nSubmit"),
        ("4. Review", "Agent → Ops"),
        ("5. Confirm\nPayment", "Collection of\nfunds"),
        ("6. Select\nPayout Bank", "Fee waterfall"),
    ]
    box_w = Inches(1.7)
    box_h = Inches(1.55)
    gap = Inches(0.28)
    start_x = Inches(0.7)
    y = Inches(2.6)

    for i, (label, sub) in enumerate(steps):
        x = start_x + i * (box_w + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shape.line.color.rgb = LINE
        tf = shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = label
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = TITLE
        p.font.name = "Calibri"
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = sub
        p2.font.size = Pt(11)
        p2.font.color.rgb = MUTED
        p2.font.name = "Calibri"
        p2.alignment = PP_ALIGN.CENTER

        if i < len(steps) - 1:
            arrow = slide.shapes.add_shape(
                MSO_AUTO_SHAPE_TYPE.RIGHT_ARROW,
                x + box_w + Inches(0.02),
                y + Inches(0.6),
                Inches(0.24),
                Inches(0.28),
            )
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT
            arrow.line.fill.background()

    add_caption(
        slide,
        "Payer submits remittance → Agent/Ops review → Ops confirms inbound funds (collection) → CapitalPay selects payout bank by fee waterfall → status and ledger stay visible.",
    )


def _card(slide, left, top, width, height, title: str, body: str) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shape.line.color.rgb = LINE
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = ACCENT
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.LEFT
    for line in body.split("\n"):
        p2 = tf.add_paragraph()
        p2.text = line
        p2.font.size = Pt(12)
        p2.font.color.rgb = TITLE
        p2.font.name = "Calibri"
        p2.alignment = PP_ALIGN.LEFT
        p2.space_before = Pt(4)


def prn_intro_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "PRN — Payment Reference Number")

    add_textbox(
        slide,
        Inches(0.6),
        Inches(1.15),
        Inches(12.1),
        Inches(0.7),
        "A 6-character code that links a bank remittance remark to a CapitalPay order, "
        "so inbound credits can be matched automatically.",
        size=15,
        color=MUTED,
    )

    cards = [
        (
            "What it is",
            "Short unique reference on the payment order\n"
            "Payer puts it in the bank transfer remark / memo\n"
            "Platform extracts PRN from the bank statement",
        ),
        (
            "Why it matters",
            "Precise match to the intended order\n"
            "Enables auto confirm when amount also matches\n"
            "Reduces manual reconciliation workload",
        ),
        (
            "Lifecycle",
            "ISSUED → BOUND → MATCHED\n"
            "or EXPIRED if unused past TTL\n"
            "Default agent-issued TTL: 72 hours",
        ),
    ]
    x0 = Inches(0.6)
    for i, (title, body) in enumerate(cards):
        _card(slide, x0 + i * Inches(4.15), Inches(2.1), Inches(3.95), Inches(3.5), title, body)

    add_caption(
        slide,
        "PRN is the bridge between off-platform bank transfers and on-platform remittance orders.",
    )


def prn_generate_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "PRN Generation")

    add_textbox(
        slide,
        Inches(0.6),
        Inches(1.15),
        Inches(12.1),
        Inches(0.45),
        "Format:  [Agent prefix]  +  [Day-of-year 001–366]  +  [Daily sequence 01–99]     e.g.  A00101",
        size=15,
        bold=True,
        color=TITLE,
    )

    _card(
        slide,
        Inches(0.6),
        Inches(1.85),
        Inches(6.0),
        Inches(3.9),
        "Structure",
        "Prefix = first character of agent_no (A–Z / 0–9)\n"
        "          or 0 when no agent is attached\n"
        "Day-of-year = local calendar day (001–366)\n"
        "Sequence = next free 01–99 for that prefix+day\n"
        "Result is always exactly 6 characters, globally unique",
    )
    _card(
        slide,
        Inches(6.85),
        Inches(1.85),
        Inches(5.85),
        Inches(3.9),
        "How it is issued",
        "1) Agent HMAC API\n"
        "   POST /api/v1/agent/prn/apply/\n"
        "   Optional: merchant_no, amount, currency,\n"
        "   reference, expire_hours\n"
        "2) Remittance order path\n"
        "   Bind an existing PRN, or auto-generate\n"
        "   when the order is created / submitted\n"
        "Payer then remits with the PRN in the remark",
    )

    add_caption(
        slide,
        "Example: agent_no AG20260001 on Jan 1 → first PRN of the day is A00101, then A00102, …",
    )


def prn_query_slide(prs) -> None:
    content_slide(
        prs,
        "Query & Match by PRN",
        "14_prn_query.png",
        "Customer Orders search by PRN (A25503) returns the matched remittance. Ops order search and bank-remark auto-match use the same code.",
    )


def closing_slide(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide)
    add_title(slide, "Scope Covered in This SOW Flow")

    bullets = [
        "Agent and Customer registration / login and KYC onboarding",
        "Remittance Fee and Agent Fee configuration",
        "Remittance apply → Agent/Ops review → Confirm payment → Select payout bank",
        "Settlement debit on the customer ledger after payout",
        "Ledger adjustment for discrepancy correction",
        "Refunds with Ops Approve / Reject",
        "Customer wallet (multi-currency demo)",
        "PRN generate, query, and bank-remark auto-match",
        "Order status Progress timeline and Account activity ledger",
    ]
    y = Inches(1.1)
    for item in bullets:
        add_textbox(slide, Inches(1.0), y, Inches(11), Inches(0.38), f"•  {item}", size=14, color=TITLE)
        y += Inches(0.42)

    add_caption(
        slide,
        "This overview covers the main process only — additional details and features are available in the full system.",
    )


def build() -> Path:
    """Full rebuild — overwrites OUT. Prefer patch_main_process_ppt for incremental edits."""
    # Pillow used for aspect ratio; install if missing
    try:
        from PIL import Image  # noqa: F401
    except ImportError:
        import subprocess
        import sys

        subprocess.check_call([sys.executable, "-m", "pip", "install", "Pillow", "-q"])

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    cover_slide(prs)
    process_map_slide(prs)

    slides = [
        ("Agent Registration & Login", "01_agent_login.png", "Agents open the Agent portal to sign in or register with a Gmail address."),
        ("Agent Registration Tab", "02_agent_register.png", "New agents create an account from the Register tab, then complete KYC onboarding."),
        ("Customer Registration & Login", "03_customer_login.png", "Customers use a separate Customer portal URL for sign-in and registration."),
        ("Customer Registration Tab", "04_customer_register.png", "Customer registration mirrors agent auth, then routes the user into onboarding."),
        ("Onboarding / KYC", "05_onboarding_kyc.png", "Profile, agent binding, financial details, and documents are submitted for Agent then Ops review."),
        ("Remittance Fee", "06_remittance_fee.png", "Ops configures the global remittance tariff: fixed fee + percent, capped by a maximum."),
        ("Agent Fee", "07_agent_fee.png", "Ops sets each agent's share of the customer remittance fee and monitors order volume."),
        ("Remittance Application", "08_remittance_apply.png", "Customers (or agents on their behalf) enter instruction and beneficiary details; quotation updates in real time."),
        ("Agent Order Review", "09_agent_orders_review.png", "Bound customer remittances first wait for agent Agree / Reject before Ops review."),
        ("Operations Order Queue", "09b_ops_orders.png", "Ops reviews approved applications; awaiting-funds orders are ready for collection."),
        ("Confirm Payment — Collection of Funds", "10b_confirm_payment.png", "Ops confirms that funds have arrived; the customer book is credited the principal. PRN auto-match can do the same step."),
        ("Select Payout Bank — Smart Routing", "10_select_payout_bank.png", "Banks are ranked by Fee Rule (lowest first); the first bank with sufficient balance is Recommended."),
        ("Status Tracking", "11_status_tracking.png", "Order detail shows charges, beneficiary data, and a Progress timeline from submit through settlement."),
        ("Account Transaction Ledger", "12_account_ledger.png", "Account activity lists CREDIT / DEBIT movements with balance after and order reference."),
        ("Ops Fund Trace (Optional)", "13_fund_trace.png", "Ops can inspect fund movement and collection balances for a remittance order."),
    ]
    for title, image, caption in slides:
        content_slide(prs, title, image, caption)

    prn_intro_slide(prs)
    prn_generate_slide(prs)
    prn_query_slide(prs)
    closing_slide(prs)

    if not OUT.exists():
        raise FileNotFoundError(
            f"Refusing to create a new PPT. Expected existing file: {OUT}"
        )
    prs.save(OUT)
    return OUT


def patch_main_process_ppt(path: Path | None = None) -> Path:
    """Insert PRN slides into the existing Main Process deck (before closing)."""
    target = path or OUT
    if not target.exists():
        raise FileNotFoundError(f"PPT not found (will not create a new one): {target}")
    prs = Presentation(str(target))

    # Idempotent: skip if PRN intro already present
    already = any(
        (getattr(sh, "text", "") or "").strip().startswith("PRN — Payment Reference Number")
        for s in prs.slides
        for sh in s.shapes
    )
    if already:
        return target

    # Update cover subtitle if present
    cover = prs.slides[0]
    for sh in cover.shapes:
        if not hasattr(sh, "text_frame"):
            continue
        text = sh.text_frame.text or ""
        if "Registration & login" in text and "PRN" not in text:
            sh.text_frame.clear()
            p = sh.text_frame.paragraphs[0]
            p.text = (
                "Registration & login · Fee configuration · Remittance · PRN · "
                "Smart payout routing · Status tracking · Account ledger"
            )
            p.font.size = Pt(16)
            p.font.color.rgb = MUTED
            p.font.name = "Calibri"

    # Update process-map caption
    if len(prs.slides) > 1:
        for sh in prs.slides[1].shapes:
            if not hasattr(sh, "text_frame"):
                continue
            text = sh.text_frame.text or ""
            if "submits remittance" in text and "PRN" not in text:
                sh.text_frame.clear()
                p = sh.text_frame.paragraphs[0]
                p.text = (
                    "Payer/Agent submits remittance request → Agent/Ops review → "
                    "PRN links bank credit to the order → CapitalPay selects payout bank "
                    "by fee/account balance waterfall → status and ledger stay visible."
                )
                p.font.size = Pt(14)
                p.font.color.rgb = MUTED
                p.font.name = "Calibri"

    # Closing is last slide — rebuild bullets to include PRN
    closing = prs.slides[-1]
    to_remove = []
    for sh in closing.shapes:
        if not hasattr(sh, "text"):
            continue
        t = sh.text or ""
        if t.startswith("•") or "main process only" in t.lower():
            to_remove.append(sh)
    for sh in to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    bullets = [
        "Agent and Customer registration / login and KYC onboarding",
        "Remittance Fee (global tariff) and Agent Fee (share of customer fee)",
        "Remittance application with live quotation",
        "Review chain: Agent Agree → Ops Approve",
        "Confirm payment: inbound funds credited to the customer book",
        "Smart routing: Select payout bank ranked by Fee Rule with balance check",
        "PRN: generate, query, and bank-remark auto-match",
        "Order status Progress timeline and Account activity ledger",
    ]
    y = Inches(1.2)
    for item in bullets:
        add_textbox(closing, Inches(1.0), y, Inches(11), Inches(0.42), f"•  {item}", size=15, color=TITLE)
        y += Inches(0.48)
    add_caption(
        closing,
        "This overview covers the main process only — additional details and features are available in the full system.",
    )

    # Append PRN slides then move closing to the end
    closing_idx = len(prs.slides) - 1
    prn_intro_slide(prs)
    prn_generate_slide(prs)
    prn_query_slide(prs)

    sld_id_lst = prs.slides._sldIdLst
    entries = list(sld_id_lst)
    closing_entry = entries[closing_idx]
    sld_id_lst.remove(closing_entry)
    sld_id_lst.append(closing_entry)

    prs.save(str(target))
    return target


def _slide_title(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame:
            text = (sh.text_frame.text or "").strip()
            if text:
                return text.splitlines()[0]
    return ""


def _rewrite_center_box(shape, heading: str, sub: str) -> None:
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = TITLE
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED
    p2.font.name = "Calibri"
    p2.alignment = PP_ALIGN.CENTER


def _set_textbox(shape, text: str, *, size=16, color=MUTED) -> None:
    tf = shape.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.name = "Calibri"


def patch_confirm_payment_slide(path: Path | None = None) -> Path:
    """Insert Confirm payment (collection) slide before Select Payout Bank."""
    target = path or OUT
    if not target.exists():
        raise FileNotFoundError(f"PPT not found (will not create a new one): {target}")
    prs = Presentation(str(target))

    already = any(_slide_title(s).startswith("Confirm Payment") for s in prs.slides)
    if already:
        return target

    cover = prs.slides[0]
    for sh in cover.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text or ""
        if "Registration & login" in text and "Confirm payment" not in text:
            _set_textbox(
                sh,
                "Registration & login · Fee configuration · Remittance · PRN · "
                "Confirm payment · Smart payout routing · Status tracking · Account ledger",
                size=16,
                color=MUTED,
            )

    if len(prs.slides) > 1:
        process = prs.slides[1]
        for sh in process.shapes:
            if not sh.has_text_frame:
                continue
            text = sh.text_frame.text or ""
            if text.startswith("5."):
                _rewrite_center_box(sh, "5. Confirm\nPayment", "Collection of\nfunds")
            elif text.startswith("6."):
                _rewrite_center_box(sh, "6. Select\nPayout Bank", "Fee waterfall")
            elif "submits remittance" in text:
                _set_textbox(
                    sh,
                    "Payer/Agent submits remittance → Agent/Ops review → Ops confirms inbound funds "
                    "(collection) → CapitalPay selects payout bank by fee waterfall → status and ledger stay visible.",
                    size=14,
                    color=MUTED,
                )

    payout_idx = None
    for i, slide in enumerate(prs.slides):
        if _slide_title(slide).startswith("Select Payout Bank"):
            payout_idx = i
            break
    if payout_idx is None:
        payout_idx = min(11, len(prs.slides) - 1)

    for i, slide in enumerate(prs.slides):
        if _slide_title(slide) == "Operations Order Queue":
            for sh in slide.shapes:
                if sh.has_text_frame and "confirms payment" in (sh.text_frame.text or ""):
                    _set_textbox(
                        sh,
                        "Ops reviews approved applications; awaiting-funds orders are ready for collection.",
                        size=13,
                        color=MUTED,
                    )

    content_slide(
        prs,
        "Confirm Payment — Collection of Funds",
        "10b_confirm_payment.png",
        "Ops confirms that funds have arrived; the customer book is credited the principal. PRN auto-match can do the same step.",
    )

    sld_id_lst = prs.slides._sldIdLst
    new_entry = list(sld_id_lst)[-1]
    sld_id_lst.remove(new_entry)
    list(sld_id_lst)[payout_idx].addprevious(new_entry)

    closing = prs.slides[-1]
    to_remove = []
    for sh in closing.shapes:
        if not hasattr(sh, "text"):
            continue
        t = sh.text or ""
        if t.startswith("•") or "main process only" in t.lower():
            to_remove.append(sh)
    for sh in to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    bullets = [
        "Agent and Customer registration / login and KYC onboarding",
        "Remittance Fee (global tariff) and Agent Fee (share of customer fee)",
        "Remittance application with live quotation",
        "Review chain: Agent Agree → Ops Approve",
        "Confirm payment: inbound funds credited to the customer book",
        "Smart routing: Select payout bank ranked by Fee Rule with balance check",
        "PRN: generate, query, and bank-remark auto-match",
        "Order status Progress timeline and Account activity ledger",
    ]
    y = Inches(1.15)
    for item in bullets:
        add_textbox(closing, Inches(1.0), y, Inches(11), Inches(0.40), f"•  {item}", size=15, color=TITLE)
        y += Inches(0.46)
    add_caption(
        closing,
        "This overview covers the main process only — additional details and features are available in the full system.",
    )

    prs.save(str(target))
    return target


def patch_prn_query_screenshot(path: Path | None = None) -> Path:
    """Replace Query & Match by PRN text cards with a live portal screenshot."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    target = path or OUT
    if not target.exists():
        raise FileNotFoundError(f"PPT not found (will not create a new one): {target}")
    image_path = SHOTS / "14_prn_query.png"
    if not image_path.exists():
        raw = SHOTS / "14_prn_customer_search.png"
        if not raw.exists():
            raise FileNotFoundError(f"Missing screenshot: {raw}")
        from PIL import Image
        with Image.open(raw) as im:
            w, h = im.size
            im.crop((0, 0, w, int(h * 0.50))).save(image_path)

    prs = Presentation(str(target))
    query_slide = None
    for slide in prs.slides:
        if _slide_title(slide).startswith("Query & Match by PRN"):
            query_slide = slide
            break
    if query_slide is None:
        raise RuntimeError("Query & Match by PRN slide not found")

    to_remove = []
    for sh in query_slide.shapes:
        is_pic = sh.shape_type == MSO_SHAPE_TYPE.PICTURE
        t = sh.text_frame.text if sh.has_text_frame else ""
        if is_pic or (
            t.startswith("Agent API")
            or t.startswith("Portal search")
            or t.startswith("Bank auto-match")
            or "API lookup" in t
            or "inbound bank-statement" in t
            or "search by PRN" in t
        ):
            to_remove.append(sh)
    for sh in to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    fit_image(query_slide, image_path)
    add_caption(
        query_slide,
        "Customer Orders search by PRN (A25503) returns the matched remittance. Ops order search and bank-remark auto-match use the same code.",
    )
    prs.save(str(target))
    return target


def patch_extra_money_slides(path: Path | None = None) -> Path:
    """Insert debit / adjustment / refund / wallet slides before the PRN section."""
    target = path or OUT
    if not target.exists():
        raise FileNotFoundError(f"PPT not found (will not create a new one): {target}")
    prs = Presentation(str(target))

    already = any(_slide_title(s).startswith("Settlement Debit") for s in prs.slides)
    if already:
        return target

    cover = prs.slides[0]
    for sh in cover.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text or ""
        if "Registration & login" in text and "Refund" not in text and "Wallet" not in text:
            _set_textbox(
                sh,
                "Registration & login · Fees · Remittance · Confirm payment · Payout routing · "
                "Ledger · Adjustment · Refund · Wallet · PRN",
                size=15,
                color=MUTED,
            )

    insert_after = None
    for i, slide in enumerate(prs.slides):
        if _slide_title(slide) == "Account Transaction Ledger":
            insert_after = i
            break
    if insert_after is None:
        for i, slide in enumerate(prs.slides):
            if _slide_title(slide).startswith("PRN — Payment Reference Number"):
                insert_after = i - 1
                break
    if insert_after is None:
        insert_after = len(prs.slides) - 2

    extras = [
        (
            "Settlement Debit — Book Deduction",
            "16_settlement_debit.png",
            "After payout, Account activity records Debit rows (Settlement payout) with balance after and order reference.",
        ),
        (
            "Ledger Adjustment — Correction",
            "17_ledger_adjustment.png",
            "Ops creates a ledger adjustment for amount or status discrepancies; Approve / Reject completes the correction.",
        ),
        (
            "Refunds",
            "18_refunds.png",
            "Ops reviews refund instructions against remittances, then credits proceeds back to the customer account.",
        ),
        (
            "Customer Wallet",
            "19_wallet.png",
            "Customer wallet shows multi-currency balances with Top up, Withdraw, Transfer, Freeze, and transaction history.",
        ),
    ]

    sld_id_lst = prs.slides._sldIdLst
    insert_at = insert_after + 1
    for title, image, caption in extras:
        content_slide(prs, title, image, caption)
        new_entry = list(sld_id_lst)[-1]
        sld_id_lst.remove(new_entry)
        list(sld_id_lst)[insert_at].addprevious(new_entry)
        insert_at += 1

    closing = prs.slides[-1]
    to_remove = []
    for sh in closing.shapes:
        if not hasattr(sh, "text"):
            continue
        t = sh.text or ""
        if t.startswith("•") or "main process only" in t.lower():
            to_remove.append(sh)
    for sh in to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    bullets = [
        "Agent and Customer registration / login and KYC onboarding",
        "Remittance Fee and Agent Fee configuration",
        "Remittance apply → Agent/Ops review → Confirm payment → Select payout bank",
        "Settlement debit on the customer ledger after payout",
        "Ledger adjustment for discrepancy correction",
        "Refunds with Ops Approve / Reject",
        "Customer wallet (multi-currency demo)",
        "PRN generate, query, and bank-remark auto-match",
        "Order status Progress timeline and Account activity ledger",
    ]
    y = Inches(1.1)
    for item in bullets:
        add_textbox(closing, Inches(1.0), y, Inches(11), Inches(0.38), f"•  {item}", size=14, color=TITLE)
        y += Inches(0.42)
    add_caption(
        closing,
        "This overview covers the main process only — additional details and features are available in the full system.",
    )

    prs.save(str(target))
    return target


def patch_prn_generation_screenshot(path: Path | None = None) -> Path:
    """Replace dense PRN Generation text cards with Agent Code edit screenshot."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    target = path or OUT
    if not target.exists():
        raise FileNotFoundError(f"PPT not found (will not create a new one): {target}")
    image_path = SHOTS / "20_agent_code_edit.png"
    if not image_path.exists():
        raise FileNotFoundError(f"Missing screenshot: {image_path}")

    prs = Presentation(str(target))
    slide = None
    for s in prs.slides:
        if _slide_title(s).startswith("PRN Generation"):
            slide = s
            break
    if slide is None:
        raise RuntimeError("PRN Generation slide not found")

    if any(sh.shape_type == MSO_SHAPE_TYPE.PICTURE for sh in slide.shapes):
        return target

    to_remove = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text or ""
        if (
            t.startswith("Structure")
            or t.startswith("How it is issued")
            or t.startswith("Format:")
            or "first PRN of the day" in t
            or "agent_no AG20260001" in t
        ):
            to_remove.append(sh)
    for sh in to_remove:
        sp = sh._element
        sp.getparent().remove(sp)

    fit_image(slide, image_path)
    add_caption(
        slide,
        "Agent Code (agent_no) is edited in Ops Profiles. Its first character becomes the PRN prefix (e.g. AG20260003 → A…).",
    )
    prs.save(str(target))
    return target


if __name__ == "__main__":
    path = patch_main_process_ppt()
    path = patch_confirm_payment_slide(path)
    path = patch_prn_query_screenshot(path)
    path = patch_extra_money_slides(path)
    path = patch_prn_generation_screenshot(path)
    print(f"Updated in place: {path}")

