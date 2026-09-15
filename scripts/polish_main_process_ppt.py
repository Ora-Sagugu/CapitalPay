"""Polish CapitalPay Main Process.pptx: order, titles, captions, process map."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "CapitalPay Main Process.pptx"
TMP = OUT.with_name("CapitalPay Main Process._polish.pptx")

TITLE = RGBColor(0x1F, 0x29, 0x37)
MUTED = RGBColor(0x5B, 0x64, 0x72)


def slide_title(slide) -> str:
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = (sh.text_frame.text or "").strip()
            if t:
                return t.splitlines()[0]
    return ""


def set_textbox(shape, text: str, *, size=14, color=MUTED, bold=False, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = "Calibri"
    p.alignment = align


def rewrite_box(shape, heading: str, sub: str):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = heading
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = TITLE
    p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER
    p2 = tf.add_paragraph()
    p2.text = sub
    p2.font.size = Pt(11)
    p2.font.color.rgb = MUTED
    p2.font.name = "Calibri"
    p2.alignment = PP_ALIGN.CENTER


def reorder(prs, desired_titles: list[str]):
    titles = [slide_title(s) for s in prs.slides]
    by_title = {}
    for i, t in enumerate(titles):
        by_title.setdefault(t, i)

    closing_title = "Scope Covered in This SOW Flow"
    used = set()
    order_idx = []
    for t in desired_titles:
        if t in by_title and t not in used:
            order_idx.append(by_title[t])
            used.add(t)
    for i, t in enumerate(titles):
        if t not in used and t != closing_title:
            order_idx.append(i)
            used.add(t)
    if closing_title in by_title:
        order_idx.append(by_title[closing_title])

    sld_id_lst = prs.slides._sldIdLst
    entries = list(sld_id_lst)
    new_entries = [entries[i] for i in order_idx]
    for e in list(sld_id_lst):
        sld_id_lst.remove(e)
    for e in new_entries:
        sld_id_lst.append(e)


def polish() -> Path:
    prs = Presentation(str(OUT))

    # Cover
    cover = prs.slides[0]
    for sh in cover.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text or ""
        if "Registration" in t or "Fees" in t or "Remittance" in t or "Confirm payment" in t:
            set_textbox(
                sh,
                "Register & login · Fees · Remittance · PRN · Confirm payment · "
                "Payout routing · Ledger · Adjustment · Refund · Wallet",
                size=15,
                color=MUTED,
            )
        if "Main Process Overview" in t:
            set_textbox(sh, "Main Process Overview", size=32, color=TITLE, bold=True)

    # Process map
    process = prs.slides[1]
    step_map = {
        "1.": ("1. Register\n& Login", "Agent /\nCustomer"),
        "2.": ("2. Fee\nConfig", "Remittance &\nAgent Fee"),
        "3.": ("3. Remittance\n& Review", "Quote → Agent\n→ Ops"),
        "4.": ("4. PRN &\nConfirm Pay", "Match inbound\nfunds"),
        "5.": ("5. Select\nPayout Bank", "Fee waterfall"),
        "6.": ("6. Ledger &\nAftercare", "Debit · Refund\n· Wallet"),
    }
    for sh in process.shapes:
        if not sh.has_text_frame:
            continue
        text = sh.text_frame.text or ""
        for prefix, (h, s) in step_map.items():
            if text.startswith(prefix):
                rewrite_box(sh, h, s)
                break
        if "submits remittance" in text or "confirms inbound" in text or "Register →" in text:
            set_textbox(
                sh,
                "Register → configure fees → apply remittance → Agent/Ops review → "
                "PRN links inbound credit → Confirm payment → Select payout bank → "
                "ledger / adjustment / refund / wallet.",
                size=13,
                color=MUTED,
            )

    # Rename PRN Generation → Agent Code — PRN Prefix
    for slide in prs.slides:
        if slide_title(slide).startswith("PRN Generation") or slide_title(slide).startswith("Agent Code"):
            for sh in slide.shapes:
                if not sh.has_text_frame:
                    continue
                t = sh.text_frame.text or ""
                if t.startswith("PRN Generation") or t.startswith("Agent Code — PRN Prefix"):
                    set_textbox(sh, "Agent Code — PRN Prefix", size=26, color=TITLE, bold=True)
                elif "Agent Code (agent_no)" in t or "Ops Profiles" in t:
                    set_textbox(
                        sh,
                        "Ops Profiles → Edit agent. Agent Code (e.g. AG20260003) supplies the PRN prefix character A.",
                        size=13,
                        color=MUTED,
                    )

    caption_updates = {
        "Agent Registration Tab": "Agent portal Register tab — create an agent account with Gmail, then complete KYC.",
        "Agent Registration & Login": "Agent portal Sign in — demo GraceNyambura@gmail.com / 123456.",
        "Customer Registration & Login": "Customer portal Sign in / Register — separate URL from the Agent portal.",
        "Settlement Debit — Book Deduction": "After payout completes, Account activity shows Debit / Settlement payout rows and running balance.",
        "Ops Fund Trace (Optional)": "Ops Fund Trace follows an order from collection VA through settlement stages.",
        "Account Transaction Ledger": "Customer Account activity: CREDIT = collection confirmed; DEBIT = settlement payout.",
        "Confirm Payment — Collection of Funds": "Ops Confirm payment credits the customer book. PRN auto-match can perform the same step.",
        "Select Payout Bank — Smart Routing": "Banks ranked by Fee Rule (lowest first); first bank with enough balance is Recommended.",
        "Query & Match by PRN": "Customer Orders search by PRN (A25503) returns the matched remittance.",
        "Ledger Adjustment — Correction": "Ops ledger adjustment corrects amount/status discrepancies (Approve / Reject).",
        "Refunds": "Ops refund queue: review remittance refunds and credit proceeds to the customer account.",
        "Customer Wallet": "Customer wallet: multi-currency balances, Top up / Withdraw / Transfer / Freeze (demo).",
    }
    for slide in prs.slides:
        title = slide_title(slide)
        if title not in caption_updates:
            continue
        text_shapes = [sh for sh in slide.shapes if sh.has_text_frame and (sh.text_frame.text or "").strip()]
        if len(text_shapes) < 2:
            continue
        bottom = max(text_shapes, key=lambda sh: sh.top)
        if (bottom.text_frame.text or "").startswith(title):
            continue
        set_textbox(bottom, caption_updates[title], size=13, color=MUTED)

    titles_now = [slide_title(s) for s in prs.slides]
    print("titles before reorder:")
    for i, t in enumerate(titles_now, 1):
        print(f"  {i:02d} {t}")

    desired = [
        titles_now[0],
        "End-to-End Main Process",
        "Agent Registration & Login",
        "Agent Registration Tab",
        "Customer Registration & Login",
        "Onboarding / KYC",
        "Remittance Fee",
        "Agent Fee",
        "Remittance Application",
        "Agent Order Review",
        "Operations Order Queue",
        "PRN — Payment Reference Number",
        "Agent Code — PRN Prefix",
        "Query & Match by PRN",
        "Confirm Payment — Collection of Funds",
        "Select Payout Bank — Smart Routing",
        "Status Tracking",
        "Account Transaction Ledger",
        "Settlement Debit — Book Deduction",
        "Ops Fund Trace (Optional)",
        "Ledger Adjustment — Correction",
        "Refunds",
        "Customer Wallet",
        "Scope Covered in This SOW Flow",
    ]
    if "Agent Code — PRN Prefix" not in titles_now and "PRN Generation" in titles_now:
        desired = [("PRN Generation" if t == "Agent Code — PRN Prefix" else t) for t in desired]

    reorder(prs, desired)

    # Closing
    closing = None
    for s in prs.slides:
        if slide_title(s).startswith("Scope Covered"):
            closing = s
            break
    if closing is not None:
        to_remove = []
        for sh in closing.shapes:
            if not hasattr(sh, "text"):
                continue
            t = sh.text or ""
            if t.startswith("•") or "main process only" in t.lower():
                to_remove.append(sh)
        for sh in to_remove:
            sh._element.getparent().remove(sh._element)

        bullets = [
            "Agent / Customer register & login, then KYC onboarding",
            "Remittance Fee and Agent Fee configuration",
            "Remittance apply → Agent review → Ops review",
            "PRN: Agent Code prefix, generate / query / bank-remark match",
            "Confirm payment (collection) → Select payout bank (smart routing)",
            "Status tracking, account ledger, and settlement debit",
            "Ledger adjustment, refunds, and customer wallet",
        ]
        y = Inches(1.25)
        for item in bullets:
            box = closing.shapes.add_textbox(Inches(1.0), y, Inches(11), Inches(0.42))
            tf = box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = f"•  {item}"
            p.font.size = Pt(15)
            p.font.color.rgb = TITLE
            p.font.name = "Calibri"
            y += Inches(0.48)
        box = closing.shapes.add_textbox(Inches(0.6), Inches(6.85), Inches(12.1), Inches(0.5))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = (
            "This overview covers the main process only — "
            "additional details and features are available in the full system."
        )
        p.font.size = Pt(13)
        p.font.color.rgb = MUTED
        p.font.name = "Calibri"

    prs.save(str(TMP))
    OUT.unlink()
    TMP.replace(OUT)

    prs2 = Presentation(str(OUT))
    print("\ntitles after polish:")
    for i, s in enumerate(prs2.slides, 1):
        print(f"  {i:02d} {slide_title(s)}")
    return OUT


if __name__ == "__main__":
    print(f"Updated: {polish()}")
